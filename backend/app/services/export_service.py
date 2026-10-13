import io
import json

from ..models import Transcript


def _format_srt_timestamp(seconds: float) -> str:
    millis = int(round(seconds * 1000))
    hours, millis = divmod(millis, 3_600_000)
    minutes, millis = divmod(millis, 60_000)
    secs, millis = divmod(millis, 1_000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"


def export_txt(transcript: Transcript) -> bytes:
    return transcript.text.encode("utf-8")


def export_srt(transcript: Transcript) -> bytes:
    segments = json.loads(transcript.segments_json or "[]")
    if not segments:
        # fall back to a single block covering the whole transcript
        segments = [{"start": 0.0, "end": max(transcript.duration_seconds or 1.0, 1.0),
                     "text": transcript.text}]

    lines = []
    for idx, seg in enumerate(segments, start=1):
        lines.append(str(idx))
        lines.append(
            f"{_format_srt_timestamp(seg['start'])} --> {_format_srt_timestamp(seg['end'])}"
        )
        lines.append(seg["text"])
        lines.append("")

    return "\n".join(lines).encode("utf-8")


def export_md(transcript: Transcript) -> bytes:
    lines = [f"# {transcript.original_name}", ""]
    if transcript.language:
        lines.append(f"*Language: {transcript.language}*")
        lines.append("")
    segments = json.loads(transcript.segments_json or "[]")
    if segments:
        for seg in segments:
            start = _format_srt_timestamp(seg["start"]).replace(",", ".")
            lines.append(f"**[{start}]** {seg['text']}")
            lines.append("")
    else:
        lines.append(transcript.text or "")
    if transcript.translated_text:
        lines.append("")
        lines.append(f"## Translation ({transcript.translated_language})")
        lines.append("")
        lines.append(transcript.translated_text)
    return "\n".join(lines).encode("utf-8")


def export_vtt(transcript: Transcript) -> bytes:
    segments = json.loads(transcript.segments_json or "[]")
    if not segments:
        segments = [{"start": 0.0, "end": max(transcript.duration_seconds or 1.0, 1.0),
                     "text": transcript.text}]

    def _vtt_timestamp(seconds: float) -> str:
        millis = int(round(seconds * 1000))
        hours, millis = divmod(millis, 3_600_000)
        minutes, millis = divmod(millis, 60_000)
        secs, millis = divmod(millis, 1_000)
        return f"{hours:02d}:{minutes:02d}:{secs:02d}.{millis:03d}"

    lines = ["WEBVTT", ""]
    for seg in segments:
        lines.append(f"{_vtt_timestamp(seg['start'])} --> {_vtt_timestamp(seg['end'])}")
        lines.append(seg["text"])
        lines.append("")
    return "\n".join(lines).encode("utf-8")


def export_docx(transcript: Transcript) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading(transcript.original_name, level=1)

    meta_parts = []
    if transcript.language:
        meta_parts.append(f"Language: {transcript.language}")
    if transcript.duration_seconds:
        meta_parts.append(f"Duration: {transcript.duration_seconds:.1f}s")
    if meta_parts:
        doc.add_paragraph(" · ".join(meta_parts)).italic = True

    for paragraph in (transcript.text or "").split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph)

    if transcript.translated_text:
        doc.add_heading(f"Translation ({transcript.translated_language})", level=2)
        for paragraph in transcript.translated_text.split("\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def export_pdf(transcript: Transcript) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import simpleSplit

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    margin = 2 * cm
    max_width = width - 2 * margin

    c.setFont("Helvetica-Bold", 14)
    c.drawString(margin, height - margin, transcript.original_name)
    c.setFont("Helvetica", 9)
    c.drawString(margin, height - margin - 16, f"Language: {transcript.language or 'auto'}")

    c.setFont("Helvetica", 11)
    y = height - margin - 40
    line_height = 14

    for paragraph in (transcript.text or "").split("\n"):
        wrapped = simpleSplit(paragraph, "Helvetica", 11, max_width) or [""]
        for line in wrapped:
            if y < margin:
                c.showPage()
                c.setFont("Helvetica", 11)
                y = height - margin
            c.drawString(margin, y, line)
            y -= line_height

    c.save()
    buffer.seek(0)
    return buffer.read()
